# -*- coding: utf-8 -*-
"""从 docs/Cornerstone项目汇报.md 生成两页汇报 PPT。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(r"d:\work\CornerstoneMock")
DOCS = ROOT / "docs"
ASSETS = DOCS / "report-assets"
PPT_OUT = DOCS / "Cornerstone项目汇报.pptx"
PPT_ALT = ROOT / "Cornerstone项目汇报简版-更新.pptx"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x2E, 0x75, 0xB6)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF5, 0xF8, 0xFC)
GREEN = RGBColor(0x70, 0xAD, 0x47)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
PLAN_GREEN = RGBColor(0x38, 0x76, 0x1D)

# --- 与 docs/Cornerstone项目汇报.md 同步 ---

SLIDE1 = {
    "title": "课题背景、目标与技术实现",
    "subtitle": "围绕 LECO 气体分析仪器远程运维场景，构建「稳定通信 + AI 诊断 + 自动处置」的智能化能力",
    "background": [
        "LECO 气体分析仪器广泛应用于工业检测场景，传统运维依赖人工现场或远程手动操作。",
        "现有模式存在效率低、响应慢、故障预判能力弱等问题。",
        "在已具备基础远程控制能力的基础上，亟需引入 AI 提升主动运维、智能诊断与高效管控水平。",
    ],
    "goals": [
        "数据通信能力：稳定运行 TCP/HTTP 远程通信、指令下发、样品登录、状态查询、网关转发与 Web 管理界面。",
        "AI 能力落地：基于仪器状态、日志、环境数据，实现异常智能检测、故障自动预警、运行趋势预测。",
        "效率提升：形成「远程控制 + 智能分析 + 自动处置」一体化运维能力。",
    ],
    "arch_intro": (
        "仪器侧（Bridge / Web / Queue / CLI → 仪器）已基本实现；"
        "扩展为各实验室部署 CornerstoneAgent，公司内网部署仪器智能体与知识库。"
    ),
    "tags": [
        ("统一协议", "TCP/XML + Bridge REST"),
        ("分层部署", "仪器侧 + 边缘 Agent + 公司内网"),
        ("REST 能力", "Bridge REST；Web /api 反代"),
        ("AI 扩展", "Agent + 智能体 + 知识库 + 大模型"),
    ],
    "summary": (
        "仪器侧通信底座已基本实现；边缘 Agent 负责采集与规则；"
        "公司内网智能体与知识库提供对话查询与知识增强。"
    ),
}

SLIDE2 = {
    "title": "项目进展及效果",
    "subtitle": "远程控制底座与多入口已落地，当前聚焦大模型对接与对话式查询",
    "stages": [
        ("阶段 0-1", "仪器侧通信底座\nBridge · Web · CLI", "✅ 已实现", GREEN),
        ("阶段 1b", "Queue 试样缓存", "✅ 已实现", GREEN),
        ("阶段 2", "边缘 CornerstoneAgent\n采集 · 规则 · 信息窗口", "⏳ 规划中", TEAL),
        ("阶段 3", "公司内网\n智能体 · 知识库 · 大模型", "⏳ 对接中", ORANGE),
        ("远期", "北向 Modbus / MQTT", "📋 暂缓", GRAY),
    ],
    "images": [
        (ASSETS / "queue.png", "Queue 悬浮窗 · 试样代码缓存"),
        (ASSETS / "web.png", "Web 分析页 · 数据与状态查询"),
    ],
    "done": [
        "Bridge / Web 分包与配置拆分，TCP 网关 + REST + XML 解析；",
        "Bridge 控制台、Web 分析页（ECharts 谱图与数据查询）；",
        "Queue 悬浮窗：试样代码手动缓存、Bridge 在线管理与发送；",
        "仪器锁定自动状态下，Web 可查询分析结果与仪器状态。",
    ],
    "current": [
        "学习了解公司大模型，争取打通网络环境；",
        "跑通基本业务逻辑：对话查询分析结果数据与仪器状态。",
    ],
    "future": [
        "探索对话功能集成到宝武聊天；",
        "仪器知识库建立：说明书、图纸、应用文档、维修经验等。",
    ],
}


def _blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    return prs.slides.add_slide(layout)


def _rect(slide, left, top, width, height, fill, *, line=None, text="", text_size=11, text_color=WHITE, bold=True):
    sh = slide.shapes.add_shape(1, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    if text:
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(text_size)
        p.font.bold = bold
        p.font.color.rgb = text_color
        p.alignment = PP_ALIGN.CENTER
    return sh


def _textbox(slide, left, top, width, height, text, *, size=13, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def _bullets(slide, left, top, width, height, title, items, *, title_color=BLUE, size=12, numbered=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = title_color
    p0.space_after = Pt(4)
    for i, item in enumerate(items, 1):
        p = tf.add_paragraph()
        p.text = f"{i}) {item}" if numbered else item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(3)
    return box


def _header(slide, title: str, subtitle: str):
    _rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.55), BLUE)
    _textbox(slide, Inches(0.45), Inches(0.08), Inches(5), Inches(0.35),
             "Cornerstone Remote Control", size=18, bold=True, color=WHITE)
    _textbox(slide, Inches(0.45), Inches(0.72), Inches(12), Inches(0.42), title, size=24, bold=True, color=BLUE)
    _textbox(slide, Inches(0.45), Inches(1.14), Inches(12.2), Inches(0.5), subtitle, size=12, color=GRAY)


def _arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = GRAY
    conn.line.width = Pt(1.5)
    return conn


def _arch_diagram(slide):
    """简版三层架构框图。"""
    top = Inches(3.55)
    _textbox(slide, Inches(0.45), Inches(3.35), Inches(12), Inches(0.25),
             "技术架构图", size=14, bold=True, color=BLUE)

    # 公司内网
    _rect(slide, Inches(0.45), top, Inches(12.4), Inches(1.05), RGBColor(0xE8, 0xF0, 0xF8), line=TEAL)
    _textbox(slide, Inches(0.55), top + Inches(0.05), Inches(1.2), Inches(0.25),
             "公司内网", size=11, bold=True, color=TEAL)
    boxes_corp = [
        (Inches(1.8), "宝武聊天\n/ 运维入口", TEAL),
        (Inches(4.5), "仪器智能体", BLUE),
        (Inches(7.2), "公司大模型", BLUE),
        (Inches(9.9), "仪器知识库", BLUE),
    ]
    bw, bh = Inches(1.85), Inches(0.72)
    by = top + Inches(0.28)
    for x, label, color in boxes_corp:
        _rect(slide, x, by, bw, bh, color, text=label, text_size=10)
    for x in (Inches(3.65), Inches(6.35), Inches(9.05)):
        _arrow(slide, x, by + bh // 2, x + Inches(0.85), by + bh // 2)

    # 各实验室
    mid = Inches(4.85)
    _rect(slide, Inches(0.45), mid, Inches(12.4), Inches(1.05), RGBColor(0xF0, 0xF8, 0xF0), line=GREEN)
    _textbox(slide, Inches(0.55), mid + Inches(0.05), Inches(1.2), Inches(0.25),
             "各实验室", size=11, bold=True, color=GREEN)
    boxes_lab = [
        (Inches(2.2), "CornerstoneAgent", TEAL),
        (Inches(5.2), "Bridge · Web · Queue · CLI\n（已实现）", GREEN),
        (Inches(8.8), "LECO 气体分析仪器", BLUE),
    ]
    by2 = mid + Inches(0.28)
    for x, label, color in boxes_lab:
        w = Inches(2.6) if "Bridge" in label else Inches(2.4)
        _rect(slide, x, by2, w, bh, color, text=label, text_size=10)
    _arrow(slide, Inches(4.6), by2 + bh // 2, Inches(5.2), by2 + bh // 2)
    _arrow(slide, Inches(7.8), by2 + bh // 2, Inches(8.8), by2 + bh // 2)

    # 公司 <-> 实验室
    _arrow(slide, Inches(5.4), by + bh, Inches(3.4), by2)
    _textbox(slide, Inches(3.8), Inches(4.55), Inches(2.2), Inches(0.35),
             "任务 · 对话 · 采集", size=9, color=GRAY, align=PP_ALIGN.CENTER)


def build_slide1(prs: Presentation):
    s = SLIDE1
    slide = _blank_slide(prs)
    _header(slide, s["title"], s["subtitle"])

    _rect(slide, Inches(0.45), Inches(1.72), Inches(6.0), Inches(1.55), LIGHT_BG, line=RGBColor(0xD0, 0xD0, 0xD0))
    _bullets(slide, Inches(0.6), Inches(1.82), Inches(5.7), Inches(1.4),
             "一、课题背景", s["background"], size=11)

    _rect(slide, Inches(6.65), Inches(1.72), Inches(6.2), Inches(1.55), LIGHT_BG, line=RGBColor(0xD0, 0xD0, 0xD0))
    _bullets(slide, Inches(6.8), Inches(1.82), Inches(5.9), Inches(1.4),
             "二、课题目标", s["goals"], size=11, numbered=True)

    _textbox(slide, Inches(0.45), Inches(3.15), Inches(12.4), Inches(0.35),
             s["arch_intro"], size=11, color=GRAY)

    _arch_diagram(slide)

    # tags
    tag_w = Inches(2.95)
    for i, (label, desc) in enumerate(s["tags"]):
        left = Inches(0.45) + i * (tag_w + Inches(0.15))
        _rect(slide, left, Inches(6.05), tag_w, Inches(0.38), TEAL if i % 2 == 0 else BLUE,
              text=label, text_size=11)
        _textbox(slide, left, Inches(6.48), tag_w, Inches(0.35), desc, size=9, color=GRAY, align=PP_ALIGN.CENTER)

    _textbox(slide, Inches(0.45), Inches(6.9), Inches(12.4), Inches(0.45),
             s["summary"], size=11, color=DARK)


def build_slide2(prs: Presentation):
    s = SLIDE2
    slide = _blank_slide(prs)
    _header(slide, s["title"], s["subtitle"])

    # 进展流程图
    _textbox(slide, Inches(0.45), Inches(1.72), Inches(3), Inches(0.3),
             "进展流程图", size=14, bold=True, color=BLUE)
    sw = Inches(2.35)
    sh = Inches(0.95)
    top = Inches(2.05)
    gap = Inches(0.12)
    x0 = Inches(0.45)
    for i, (title, body, status, color) in enumerate(s["stages"]):
        left = x0 + i * (sw + gap)
        _rect(slide, left, top, sw, sh, color,
              text=f"{title}\n{body}\n{status}", text_size=9)
        if i < len(s["stages"]) - 1:
            _arrow(slide, left + sw, top + sh // 2, left + sw + gap, top + sh // 2)

    # 两张已完成截图
    img_top = Inches(3.15)
    img_h = Inches(1.55)
    img_w = Inches(5.95)
    for i, (path, caption) in enumerate(s["images"]):
        left = Inches(0.45) + i * (img_w + Inches(0.35))
        _rect(slide, left, img_top, img_w, img_h + Inches(0.35), LIGHT_BG, line=RGBColor(0xD0, 0xD0, 0xD0))
        slide.shapes.add_picture(str(path), left + Inches(0.08), img_top + Inches(0.08),
                                 img_w - Inches(0.16), img_h)
        _textbox(slide, left, img_top + img_h + Inches(0.1), img_w, Inches(0.28),
                 caption, size=10, color=TEAL, align=PP_ALIGN.CENTER)

    # 三列文字
    col_top = Inches(5.15)
    col_h = Inches(2.1)
    col_w = Inches(3.95)
    panels = [
        ("已完成 · 基础能力", s["done"], LIGHT_BG, BLUE),
        ("当前目标", s["current"], RGBColor(0xE8, 0xF5, 0xE9), PLAN_GREEN),
        ("后续计划", s["future"], RGBColor(0xFF, 0xF3, 0xE8), ORANGE),
    ]
    for i, (title, items, bg, tc) in enumerate(panels):
        left = Inches(0.45) + i * (col_w + Inches(0.22))
        _rect(slide, left, col_top, col_w, col_h, bg, line=RGBColor(0xD0, 0xD0, 0xD0))
        _bullets(slide, left + Inches(0.12), col_top + Inches(0.1),
                 col_w - Inches(0.24), col_h - Inches(0.15),
                 title, items, title_color=tc, size=10)


def main():
    for path, _ in SLIDE2["images"]:
        if not path.exists():
            raise FileNotFoundError(path)

    ref = DOCS / "Cornerstone项目汇报简版.pptx"
    if ref.exists():
        prs = Presentation()
        old = Presentation(str(ref))
        prs.slide_width = old.slide_width
        prs.slide_height = old.slide_height
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    build_slide1(prs)
    build_slide2(prs)

    prs.save(str(PPT_OUT))
    prs.save(str(PPT_ALT))
    print(f"Saved: {PPT_OUT}")
    print(f"Saved: {PPT_ALT}")


if __name__ == "__main__":
    main()
